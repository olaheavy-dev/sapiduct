"""Small files built in the tests, so no binary fixture is committed."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from pypdf import PdfReader, PdfWriter


def pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page ('' makes a blank page)."""
    objects: list[bytes] = []
    first_page = 4
    kids = " ".join(f"{first_page + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        content_id = first_page + 2 * i + 1
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def encrypted_pdf() -> bytes:
    """The sample PDF, password protected (RC4, which pypdf can write without extra packages)."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf(["secret page"]))))
    writer.encrypt("owner-password", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx() -> bytes:
    document = Document()
    document.add_heading("Thermodynamics", level=1)
    document.add_paragraph("Energy is conserved.")
    document.add_heading("Entropy", level=2)
    document.add_paragraph("Entropy never decreases in an isolated system.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Law"
    table.rows[0].cells[1].text = "Statement"
    table.rows[1].cells[0].text = "Second"
    table.rows[1].cells[1].text = "Entropy rises"
    document.add_heading("Heat engines", level=1)
    document.add_paragraph("No engine is perfectly efficient.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def zip_bomb(declared_mb: int) -> bytes:
    """A real zip whose single member expands to `declared_mb` megabytes of zeros."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (declared_mb * 1024 * 1024))
    return out.getvalue()


VTT = """WEBVTT

00:00:01.000 --> 00:00:04.000
Welcome to the lecture on entropy.

00:00:04.500 --> 00:00:09.000
Today we cover the second law.

00:01:10.000 --> 00:01:14.000
Heat flows from hot to cold.
"""

MARKDOWN = """Intro text before any heading.

# Physics

Overview of physics.

## Entropy

Entropy measures disorder.

```
# not a heading, it is code
```

## Heat

Heat flows downhill.

# Chemistry

Atoms combine.
"""

HTML = """<!doctype html><html><head><title>Notes</title><style>p{color:red}</style>
<script>alert('x')</script></head><body>
<nav>Home | About | Cookies</nav>
<h1>Physics</h1><p>Overview &amp; scope.</p>
<h2>Entropy</h2><p>Entropy measures <b>disorder</b>.</p><ul><li>one</li><li>two</li></ul>
<h1>Chemistry</h1><p>Atoms combine.</p>
</body></html>"""

TEXT = """First paragraph line one.
First paragraph line two.

Second paragraph.


Third paragraph.
"""
